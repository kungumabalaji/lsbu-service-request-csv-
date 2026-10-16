"""
Three test scenarios for the triage prototype.

  Test 1 - the 10 requests supplied with the task are processed correctly
  Test 2 - missing and invalid data is caught and never acknowledged
  Test 3 - edge cases: no keyword match, category disagreement, mixed case,
           completed requests and spreadsheet formula injection

Run with:  python -m pytest -v
"""
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import triage

ROOT = Path(__file__).resolve().parent.parent
CATEGORY_RULES, REVIEW_RULES = triage.load_rules(ROOT / "rules.csv")
TODAY = date(2026, 10, 15)          # fixed date so results never change


def run(csv_name: str) -> list[dict]:
    requests = triage.read_requests(ROOT / "data" / csv_name)
    return triage.triage(requests, CATEGORY_RULES, REVIEW_RULES, today=TODAY)


def test_1_supplied_requests_are_processed_correctly(tmp_path):
    results = run("requests.csv")
    by_id = {r["Request ID"]: r for r in results}

    # every row is read and passes validation
    assert len(results) == 10
    assert all(r["Validation Status"] == "Valid" for r in results)

    # keyword rules agree with the submitted Request Type for every row
    assert all(r["Assigned Category"] == r["Request Type"] for r in results)

    # exactly the expected requests are flagged, for the right reasons
    flagged = {r["Request ID"] for r in results if r["Needs Human Review"] == "Yes"}
    assert flagged == {"REQ-001", "REQ-002", "REQ-004", "REQ-006", "REQ-007", "REQ-009"}
    assert "Finance-related" in by_id["REQ-004"]["Review Reasons"]
    assert "External access" in by_id["REQ-007"]["Review Reasons"]
    assert "Access request" in by_id["REQ-001"]["Review Reasons"]

    # REQ-001 mentions access AND dashboard/reporting: first rule wins, ambiguity recorded
    assert by_id["REQ-001"]["Assigned Category"] == "Access & Permissions"
    assert by_id["REQ-001"]["Also Matched"] == "Reporting & Analytics"
    assert by_id["REQ-001"]["Category Confidence"] == "Medium"

    # acknowledgements drafted for open requests, not for the completed one
    assert by_id["REQ-003"]["Suggested Acknowledgement"].startswith(
        "Thank you for your request REQ-003")
    assert by_id["REQ-010"]["Suggested Acknowledgement"].startswith("Not required")

    # the Excel report is created with all three sheets
    report = tmp_path / "report.xlsx"
    summary = triage.build_summary(results, CATEGORY_RULES)
    triage.write_report(results, summary, CATEGORY_RULES, REVIEW_RULES, report)
    assert load_workbook(report).sheetnames == ["Results", "Summary", "Rules"]


@pytest.mark.parametrize("row, expected_issue", [
    (0, "Missing Request Summary"),
    (1, "Priority 'Urgent' is not one of"),
    (2, "not a valid DD/MM/YYYY date"),           # 31/02/2026 does not exist
    (3, "Duplicate Request ID"),                  # REQ-103 appears twice
    (4, "Request ID must look like REQ-001"),     # REQ104 has no hyphen
    (5, "Submitted Date is in the future"),
    (11, "Missing Requester Role"),               # only spaces in the cell
])
def test_2_missing_and_invalid_data_is_caught(row, expected_issue):
    result = run("test_requests.csv")[row]

    assert result["Validation Status"] == "Invalid"
    assert expected_issue in result["Validation Issues"]
    assert result["Needs Human Review"] == "Yes"
    assert result["Suggested Acknowledgement"].startswith("Not drafted")


def test_3_edge_cases_are_handled_safely(tmp_path):
    by_id = {r["Request ID"]: r for r in run("test_requests.csv")}

    # no keyword matches: not guessed, sent to a person
    assert by_id["REQ-106"]["Assigned Category"] == "Uncategorised"
    assert "No category rule matched" in by_id["REQ-106"]["Review Reasons"]

    # rules disagree with the Request Type the requester chose
    assert by_id["REQ-107"]["Assigned Category"] == "Reporting & Analytics"
    assert by_id["REQ-107"]["Category Confidence"] == "Low"
    assert "Category unclear" in by_id["REQ-107"]["Review Reasons"]
    assert "logged it as" not in by_id["REQ-107"]["Suggested Acknowledgement"]

    # mixed case and extra spaces are tidied; completed work is not acknowledged
    assert by_id["REQ-109"]["Priority"] == "Medium"
    assert by_id["REQ-109"]["Channel"] == "Email"
    assert by_id["REQ-109"]["Validation Status"] == "Valid"
    assert by_id["REQ-109"]["Suggested Acknowledgement"].startswith("Not required")

    # UK dates: 13/10/2026 is valid (it would fail if read as US month/day)
    assert by_id["REQ-110"]["Validation Status"] == "Valid"

    # several review reasons are all recorded
    assert "High priority" in by_id["REQ-110"]["Review Reasons"]
    assert "External access" in by_id["REQ-110"]["Review Reasons"]

    # formula injection: flagged, and stored as plain text in Excel
    assert "formula character" in by_id["REQ-108"]["Review Reasons"]
    results = list(by_id.values())
    report = tmp_path / "report.xlsx"
    triage.write_report(results, triage.build_summary(results, CATEGORY_RULES),
                        CATEGORY_RULES, REVIEW_RULES, report)
    sheet = load_workbook(report)["Results"]
    cell = next(c for c in sheet["E"] if c.value and "SUM(1+1)" in str(c.value))
    assert cell.data_type == "s" and cell.value.startswith("'=")   # text, not a formula

    # a file with a missing column is rejected with a clear message
    broken = tmp_path / "broken.csv"
    broken.write_text("Request ID,Submitted Date,Request Summary\nREQ-001,01/10/2026,Test\n")
    with pytest.raises(ValueError, match="missing required column"):
        triage.read_requests(broken)
