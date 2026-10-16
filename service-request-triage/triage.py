#!/usr/bin/env python3
"""
Service Request Triage - prototype

Reads service requests from a CSV file and, for each request:
  1. checks that the required information is present and valid
  2. assigns a category using simple keyword rules (from rules.csv)
  3. drafts a suggested acknowledgement (it is never sent automatically)
  4. decides whether a person needs to review it, and records why
It then writes an Excel report (Results, Summary, Rules) and prints a summary.

Usage:
    python triage.py data/requests.csv output/triage_output.xlsx
    python triage.py data/requests.csv output/triage_output.xlsx --rules rules.csv
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# ---------------------------------------------------------------------------
# Configuration: business rules that can change without touching the logic.
# Category and review keywords live in rules.csv so the team can edit them
# in Excel without a developer.
# ---------------------------------------------------------------------------
REQUIRED_COLUMNS = [
    "Request ID", "Submitted Date", "Requester Role", "Request Type",
    "Request Summary", "Priority", "Channel", "Status",
]
ALLOWED_PRIORITIES = ["Low", "Medium", "High"]
ALLOWED_CHANNELS = ["Service Portal", "Email", "Teams"]
ALLOWED_STATUSES = ["New", "In Review", "Assigned", "Planned", "In Progress",
                    "On Hold", "Completed", "Closed", "Cancelled"]
CLOSED_STATUSES = {"Completed", "Closed", "Cancelled"}
ALWAYS_REVIEW_CATEGORIES = {"Access & Permissions": "Access request - approval required"}
ID_PATTERN = re.compile(r"^REQ-\d{3,}$")
DATE_FORMAT = "%d/%m/%Y"            # UK format: 01/10/2026 means 1 October 2026
MIN_SUMMARY_LENGTH = 10
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")   # spreadsheet formula injection
UNCATEGORISED = "Uncategorised"

DEFAULT_RULES = Path(__file__).with_name("rules.csv")


# ---------------------------------------------------------------------------
# Step 1: read the inputs
# ---------------------------------------------------------------------------
def load_rules(path: str | Path) -> tuple[list, list]:
    """Read rules.csv. Category rules are checked in file order (first match wins)."""
    category_rules, review_rules = [], []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for line_no, row in enumerate(csv.DictReader(f), start=2):
            rule_type = (row.get("rule_type") or "").strip().lower()
            name = (row.get("name") or "").strip()
            keywords = [k.strip().lower() for k in (row.get("keywords") or "").split(";") if k.strip()]
            if not name or not keywords:
                raise ValueError(f"{path}: line {line_no} needs a name and at least one keyword")
            if rule_type == "category":
                category_rules.append((name, keywords))
            elif rule_type == "review":
                review_rules.append((name, keywords))
            else:
                raise ValueError(f"{path}: line {line_no} has unknown rule_type '{rule_type}'")
    if not category_rules:
        raise ValueError(f"{path}: no category rules found")
    return category_rules, review_rules


def read_requests(path: str | Path) -> list[dict]:
    """Read the requests CSV as plain text, trimming spaces and skipping blank lines."""
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = [h.strip() for h in (reader.fieldnames or [])]
        missing = [c for c in REQUIRED_COLUMNS if c not in headers]
        if missing:
            raise ValueError(f"{path}: missing required column(s): {', '.join(missing)}")
        reader.fieldnames = headers
        requests = []
        for row in reader:
            clean = {k: (v or "").strip() for k, v in row.items() if k is not None}
            if any(clean.values()):
                requests.append(clean)
    return requests


# ---------------------------------------------------------------------------
# Step 2: check that required information is present and valid
# ---------------------------------------------------------------------------
def normalise(value: str, allowed: list[str]) -> str | None:
    """Match a value to an allowed option ignoring case ('high' -> 'High')."""
    return {a.lower(): a for a in allowed}.get(value.strip().lower())


def validate(req: dict, id_counts: Counter, valid_types: list[str], today: date) -> list[str]:
    """Return every problem found in one request (an empty list means valid)."""
    issues = [f"Missing {col}" for col in REQUIRED_COLUMNS if not req[col]]

    request_id = req["Request ID"]
    if request_id and not ID_PATTERN.match(request_id):
        issues.append("Request ID must look like REQ-001")
    if request_id and id_counts[request_id] > 1:
        issues.append(f"Duplicate Request ID (appears {id_counts[request_id]} times)")

    if req["Submitted Date"]:
        try:
            submitted = datetime.strptime(req["Submitted Date"], DATE_FORMAT).date()
        except ValueError:
            issues.append("Submitted Date is not a valid DD/MM/YYYY date")
        else:
            if submitted > today:
                issues.append("Submitted Date is in the future")

    for col, allowed in (("Priority", ALLOWED_PRIORITIES), ("Channel", ALLOWED_CHANNELS),
                         ("Status", ALLOWED_STATUSES), ("Request Type", valid_types)):
        if req[col] and normalise(req[col], allowed) is None:
            issues.append(f"{col} '{req[col]}' is not one of: {', '.join(allowed)}")

    if req["Request Summary"] and len(req["Request Summary"]) < MIN_SUMMARY_LENGTH:
        issues.append("Request Summary is too short to understand")
    return issues


# ---------------------------------------------------------------------------
# Step 3: categorise with simple keyword rules
# ---------------------------------------------------------------------------
def keyword_hits(text: str, keywords: list[str]) -> list[str]:
    """Keywords found in the text as whole words, ignoring case."""
    text = text.lower()
    return [kw for kw in keywords if re.search(rf"(?<!\w){re.escape(kw)}(?!\w)", text)]


def categorise(summary: str, category_rules: list) -> tuple[str, list[str], list[str]]:
    """Return (category, matched keywords, other categories that also matched)."""
    matches = [(name, hits) for name, kws in category_rules if (hits := keyword_hits(summary, kws))]
    if not matches:
        return UNCATEGORISED, [], []
    (category, hits), others = matches[0], [name for name, _ in matches[1:]]
    return category, hits, others


def confidence(category: str, other_matches: list[str], request_type: str) -> str:
    """High: one clear match that agrees with the submitted Request Type.
    Medium: agrees, but other categories matched too. Low: no match or disagreement."""
    if category == UNCATEGORISED or category != request_type:
        return "Low"
    return "Medium" if other_matches else "High"


# ---------------------------------------------------------------------------
# Step 4: draft a suggested acknowledgement (a person checks and sends it)
# ---------------------------------------------------------------------------
def draft_acknowledgement(req: dict, category: str, conf: str, issues: list[str]) -> str:
    if issues:
        return "Not drafted - fix the data issues first, then acknowledge."
    if req["Status"] in CLOSED_STATUSES:
        return f"Not required - request is already {req['Status'].lower()}."
    logged_as = f" as {category}" if conf != "Low" else ""
    next_step = ("Because it is marked High priority, it has been passed to the team for prompt review."
                 if req["Priority"] == "High"
                 else "It has been added to the team's queue for review.")
    return (f"Thank you for your request {req['Request ID']}. We have logged it{logged_as} "
            f"with {req['Priority']} priority. {next_step} We will contact you with next steps.")


# ---------------------------------------------------------------------------
# Step 5: decide whether a person needs to review it, and why
# ---------------------------------------------------------------------------
def review_reasons(req: dict, issues: list[str], category: str, conf: str,
                   review_rules: list) -> list[str]:
    reasons = []
    if issues:
        reasons.append("Data issues - correct before processing")
    if req["Request Summary"].startswith(FORMULA_PREFIXES):
        reasons.append("Summary starts with a formula character - neutralised, check content")
    if req["Status"] in CLOSED_STATUSES:
        return reasons                      # finished work needs no triage

    if req["Priority"] == "High":
        reasons.append("High priority - prompt review by team lead")
    if category in ALWAYS_REVIEW_CATEGORIES:
        reasons.append(ALWAYS_REVIEW_CATEGORIES[category])
    for name, keywords in review_rules:
        if keyword_hits(req["Request Summary"], keywords):
            reasons.append(name)
    if category == UNCATEGORISED:
        reasons.append("No category rule matched - categorise manually")
    elif conf == "Low":
        reasons.append(f"Category unclear - rules suggest '{category}' but Request Type "
                       f"is '{req['Request Type'] or 'blank'}'")
    return reasons


# ---------------------------------------------------------------------------
# Run all steps for every request
# ---------------------------------------------------------------------------
def triage(requests: list[dict], category_rules: list, review_rules: list,
           today: date | None = None) -> list[dict]:
    today = today or date.today()
    valid_types = [name for name, _ in category_rules]
    controlled = (("Priority", ALLOWED_PRIORITIES), ("Channel", ALLOWED_CHANNELS),
                  ("Status", ALLOWED_STATUSES), ("Request Type", valid_types))
    id_counts = Counter(r["Request ID"] for r in requests if r["Request ID"])

    results = []
    for original in requests:
        req = dict(original)
        for col, allowed in controlled:
            req[col] = normalise(req[col], allowed) or req[col]

        issues = validate(req, id_counts, valid_types, today)
        category, hits, others = categorise(req["Request Summary"], category_rules)
        conf = confidence(category, others, req["Request Type"])
        reasons = review_reasons(req, issues, category, conf, review_rules)

        results.append({
            **{col: req[col] for col in REQUIRED_COLUMNS},
            "Validation Status": "Invalid" if issues else "Valid",
            "Validation Issues": "; ".join(issues),
            "Assigned Category": category,
            "Matched Keywords": ", ".join(hits),
            "Also Matched": ", ".join(others),
            "Category Confidence": conf,
            "Suggested Acknowledgement": draft_acknowledgement(req, category, conf, issues),
            "Needs Human Review": "Yes" if reasons else "No",
            "Review Reasons": "; ".join(reasons),
        })
    return results


# ---------------------------------------------------------------------------
# Step 6: summary and output file
# ---------------------------------------------------------------------------
def build_summary(results: list[dict], category_rules: list) -> list[tuple[str, object]]:
    total = len(results)
    invalid = sum(r["Validation Status"] == "Invalid" for r in results)
    flagged = [r["Request ID"] or "(no ID)" for r in results if r["Needs Human Review"] == "Yes"]
    drafted = sum(r["Suggested Acknowledgement"].startswith("Thank you") for r in results)

    rows = [
        ("Total requests processed", total),
        ("Valid", total - invalid),
        ("Invalid (data issues)", invalid),
        ("Need human review", len(flagged)),
        ("Requests to review", ", ".join(flagged) or "None"),
        ("No manual triage needed", total - len(flagged)),
        ("Acknowledgements drafted", drafted),
    ]
    categories = Counter(r["Assigned Category"] for r in results)
    for name in [n for n, _ in category_rules] + [UNCATEGORISED]:
        rows.append((f"Category: {name}", categories.get(name, 0)))
    priorities = Counter(r["Priority"] for r in results)
    for p in ALLOWED_PRIORITIES:
        rows.append((f"Priority: {p}", priorities.get(p, 0)))
    reason_counts = Counter(reason.split(" - ")[0]
                            for r in results if r["Review Reasons"]
                            for reason in r["Review Reasons"].split("; "))
    for reason, count in reason_counts.most_common():
        rows.append((f"Review reason: {reason}", count))
    return rows


def make_safe(value):
    """Stop spreadsheet formula injection: text starting with = + - @ is stored
    as plain text (prefixed with an apostrophe) so Excel never runs it."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")    # amber: needs human review
INVALID_FILL = PatternFill("solid", fgColor="F8CBAD")   # red: data issues
WIDE_COLUMNS = {"Request Summary": 48, "Validation Issues": 40, "Suggested Acknowledgement": 60,
                "Review Reasons": 48, "Requester Role": 24, "Request Type": 22,
                "Assigned Category": 22, "Also Matched": 22, "Matched Keywords": 18}


def _style_sheet(ws, widths: dict, default_width: int = 15) -> None:
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = widths.get(cell.value, default_width)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def write_report(results: list[dict], summary: list, category_rules: list,
                 review_rules: list, path: str | Path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Results"
    headers = list(results[0].keys()) if results else REQUIRED_COLUMNS
    ws.append(headers)
    for r in results:
        ws.append([make_safe(r[h]) for h in headers])
        fill = (INVALID_FILL if r["Validation Status"] == "Invalid"
                else REVIEW_FILL if r["Needs Human Review"] == "Yes" else None)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    _style_sheet(ws, WIDE_COLUMNS)
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Summary")
    ws.append(["Item", "Value"])
    for item, value in summary:
        ws.append([item, make_safe(value)])
    _style_sheet(ws, {"Item": 42, "Value": 60})

    ws = wb.create_sheet("Rules")
    ws.append(["Rule type", "Name", "Keywords / condition"])
    for name, keywords in category_rules:
        ws.append(["Category (first match wins)", name, ", ".join(keywords)])
    for name, keywords in review_rules:
        ws.append(["Human review (keywords)", name, ", ".join(keywords)])
    built_in = [
        ("Data issues", "Any required field missing or invalid"),
        ("High priority", "Priority = High"),
        *[(reason, f"Category = {cat}") for cat, reason in ALWAYS_REVIEW_CATEGORIES.items()],
        ("No category rule matched", "No keyword found in the Request Summary"),
        ("Confirm category", "Rule category differs from the submitted Request Type"),
        ("Formula character", "Summary starts with = + - @ (spreadsheet injection risk)"),
        ("Closed requests", "Completed / Closed / Cancelled: no acknowledgement, no triage"),
    ]
    for name, condition in built_in:
        ws.append(["Human review (built-in)", name, condition])
    _style_sheet(ws, {"Rule type": 28, "Name": 52, "Keywords / condition": 70})

    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Triage service requests from a CSV file.")
    parser.add_argument("input_csv", help="CSV file of service requests")
    parser.add_argument("output_xlsx", nargs="?", default="output/triage_output.xlsx",
                        help="Excel report to create (default: output/triage_output.xlsx)")
    parser.add_argument("--rules", default=DEFAULT_RULES, help="rules file (default: rules.csv)")
    args = parser.parse_args(argv)

    try:
        category_rules, review_rules = load_rules(args.rules)
        requests = read_requests(args.input_csv)
    except (OSError, ValueError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1

    results = triage(requests, category_rules, review_rules)
    summary = build_summary(results, category_rules)

    output = Path(args.output_xlsx)
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        write_report(results, summary, category_rules, review_rules, output)
    except PermissionError:
        print(f"Error: cannot write {output}. Close it in Excel and try again.", file=sys.stderr)
        return 1

    print("SERVICE REQUEST TRIAGE SUMMARY")
    print("=" * 60)
    for item, value in summary:
        print(f"{item:<42} {value}")
    print("=" * 60)
    print(f"Report saved to {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
